import logging
import docx
from .base import BaseDocumentParser

logger = logging.getLogger('enterprise')

class DOCXParser(BaseDocumentParser):
    """
    Concrete parser extracting paragraphs, headings, and preserving tabular row/column structure
    from Microsoft Word (.docx) documents.
    """
    def parse(self, file_path):
        try:
            doc = docx.Document(file_path)
            paragraphs = []
            headings = []
            tables_data = []

            for p in doc.paragraphs:
                txt = p.text.strip()
                if not txt:
                    continue
                if p.style.name.startswith('Heading'):
                    headings.append(txt)
                    paragraphs.append(f"## {txt}")
                else:
                    paragraphs.append(txt)

            for t_idx, table in enumerate(doc.tables, 1):
                t_rows = []
                for row in table.rows:
                    r_cells = [cell.text.strip() for cell in row.cells]
                    if any(r_cells):
                        t_rows.append(r_cells)
                if t_rows:
                    tables_data.append({
                        "table_index": t_idx,
                        "rows": t_rows
                    })
                    table_str = "\n".join([" | ".join(r) for r in t_rows])
                    paragraphs.append(f"\n[Table {t_idx}]\n{table_str}\n")

            full_text = "\n\n".join(paragraphs)

            return {
                "content": full_text,
                "structured_data": {
                    "headings": headings,
                    "tables": tables_data,
                    "paragraph_count": len(doc.paragraphs),
                    "table_count": len(doc.tables)
                },
                "metadata": {
                    "paragraph_count": len(doc.paragraphs),
                    "table_count": len(doc.tables)
                },
                "parser_type": "DOCX",
                "processing_status": "PARSED"
            }
        except Exception as e:
            logger.error(f"Failed parsing DOCX file at {file_path}: {str(e)}", exc_info=True)
            raise e

DocxParser = DOCXParser
